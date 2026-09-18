# -*- coding: utf-8 -*-
"""Genera el kit de la prueba de campo (R1): tablero de calibración A4 a escala 1:1
y la planilla de registro con las tomas planificadas.

Uso: python r1_make_fieldkit.py
Salida: OBSERVACIONES/RESPUESTA/protocolo_campo/
"""
import os
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE = r"C:\Users\José Luis\Desktop\1. MAESTRIA\CICLO II\DEEP LEARNING\Labelme"
OUT = os.path.join(BASE, "OBSERVACIONES", "RESPUESTA", "protocolo_campo")
os.makedirs(OUT, exist_ok=True)

# ---------------------------------------------------------------- tablero
SQ = 25.0            # lado del cuadro en mm
NX, NY = 10, 7       # cuadros -> 9 x 6 esquinas internas
W, H = 297.0, 210.0  # A4 horizontal en mm

fig = plt.figure(figsize=(W / 25.4, H / 25.4))
ax = fig.add_axes([0, 0, 1, 1])
ax.set_xlim(0, W)
ax.set_ylim(0, H)
ax.set_aspect("equal")
ax.axis("off")
x0 = (W - NX * SQ) / 2
y0 = 24.0
# una sola imagen binaria con extensión exacta: bordes nítidos, sin costuras de antialiasing
board = [[1 - ((i + j) % 2) for i in range(NX)] for j in range(NY)]
ax.imshow(board, cmap="gray_r", vmin=0, vmax=1, interpolation="nearest",
          extent=(x0, x0 + NX * SQ, y0, y0 + NY * SQ), origin="lower")
ax.set_xlim(0, W)  # imshow reajusta los límites: se restauran a la hoja completa
ax.set_ylim(0, H)
# barra de control de escala: 100 mm exactos
bx, by = x0, 10.0
ax.plot([bx, bx + 100], [by, by], color="black", lw=1.2)
for t in range(0, 101, 10):
    ax.plot([bx + t, bx + t], [by, by + (3 if t % 50 else 5)], color="black", lw=0.8)
ax.text(bx + 104, by, "100 mm: verifique con una regla", va="center", fontsize=9)
ax.text(W - x0, by, "RICUY R1. Tablero 9 x 6 esquinas internas, cuadro de 25 mm.\n"
        "Imprimir al 100 % (tamaño real, sin 'ajustar a la página').",
        va="center", ha="right", fontsize=8)
fig.savefig(os.path.join(OUT, "tablero_calibracion_A4.pdf"))
plt.close(fig)

# ---------------------------------------------------------------- planilla
wb = Workbook()
hdr_font = Font(bold=True, color="FFFFFF")
hdr_fill = PatternFill("solid", fgColor="4F6D7A")
fill_in = PatternFill("solid", fgColor="FFF2CC")   # celdas a completar


def header(ws, cols, widths):
    for c, (name, w) in enumerate(zip(cols, widths), start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"


ws = wb.active
ws.title = "Montaje"
header(ws, ["Dato", "Valor", "Unidad / cómo se mide"], [34, 22, 60])
rows = [
    ("Fecha", "", "dd/mm/aaaa"),
    ("Lugar (dirección o referencia)", "", ""),
    ("Marca y modelo del celular", "", "ej. Samsung A54"),
    ("Resolución del video", "1920x1080", "confirmar en la app de cámara"),
    ("FPS", "30", ""),
    ("Lente usado", "1x principal", "no usar 0.5x ni 2x"),
    ("Estabilización de video apagada", "", "sí / no"),
    ("HDR y encuadre automático apagados", "", "sí / no"),
    ("h_cam: altura del lente sobre el suelo", "", "cm, con cinta vertical desde el suelo al centro del lente"),
    ("Inclinación de la cámara", "", "grados; + hacia arriba, - hacia abajo (app de nivel)"),
    ("Distancia del lente al borde delantero de la mototaxi", "", "cm (referencia)"),
    ("Hora de inicio", "", "hh:mm"),
    ("Clima / luz", "", "nublado, sol, sombra..."),
    ("Nombre de quien anota", "", ""),
]
for r, row in enumerate(rows, start=2):
    for c, v in enumerate(row, start=1):
        cell = ws.cell(row=r, column=c, value=v)
        if c == 2:
            cell.fill = fill_in

ws2 = wb.create_sheet("Tomas")
cols = ["toma", "tipo", "objetivo", "altura real (m)", "distancia sobre el eje (m)",
        "desplazamiento lateral (m)", "archivo de video", "hora", "notas"]
header(ws2, cols, [7, 14, 16, 14, 16, 16, 26, 9, 36])
plan = [("calibración", "tablero", "", "", "", "30 a 40 s moviendo el tablero a 1 a 3 m"),
        ("calibración", "marca en el suelo", "", 8.0, 0.0, "cruz de cinta a 8.00 m, sin personas, 5 s")]
for p in (1, 2, 3):
    for d in (3, 4, 5, 7, 10, 15, 20):
        plan.append(("persona", f"Persona {p}", "", d, 0.0, ""))
for d in (5, 10):
    plan.append(("lateral", "Persona 1", "", d, 1.5, "1.5 m a la derecha, escuadra 3-4-5"))
for d in (5, 10, 15):
    plan.append(("vehículo", "Auto", "", d, 0.0, "distancia al parachoques trasero"))
for d in (5, 10):
    plan.append(("vehículo", "Mototaxi", "", d, 0.0, "distancia a la parte trasera"))
plan.append(("opcional", "Reductor o paso peatonal", "", "", 0.0, "distancia al borde más cercano"))
plan.append(("naturalista", "Recorrido 5 a 10 min", "", "", "", "no se usa para entrenar"))
for r, (tipo, obj, alt, dist, lat, nota) in enumerate(plan, start=2):
    vals = [r - 1, tipo, obj, alt, dist, lat, "", "", nota]
    for c, v in enumerate(vals, start=1):
        cell = ws2.cell(row=r, column=c, value=v)
        if c in (4, 7, 8) or (c == 5 and v == ""):
            cell.fill = fill_in

ws3 = wb.create_sheet("Estaturas")
header(ws3, ["Persona", "Nombre", "Estatura con zapatos (m)"], [12, 26, 24])
for r, p in enumerate(("Persona 1", "Persona 2", "Persona 3"), start=2):
    ws3.cell(row=r, column=1, value=p)
    for c in (2, 3):
        ws3.cell(row=r, column=c).fill = fill_in
wb.save(os.path.join(OUT, "registro_campo.xlsx"))
print("OK ->", OUT, "|", len(plan), "tomas planificadas")
